#!/usr/bin/env python3
"""Korespondencja seryjna. Szablon DOCX z polami {imie}, {adres}... + arkusz
XLSX (pierwszy wiersz = nazwy pol) -> osobny DOCX (opcjonalnie PDF) dla
kazdego wiersza. Formatowanie szablonu zostaje zachowane.

Uruchomienie: python korespondencja.py            (GUI)
              python korespondencja.py --selftest (test logiki)
"""
import datetime
import os
import re
import sys
import threading
import traceback

if getattr(sys, "frozen", False):
    APP_DIR = os.path.dirname(sys.executable)
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))

FIELD = re.compile(r"\{([^{}\n]+)\}")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

# ------------------------------------------------------------------ dane ----


def key(name):
    return str(name).strip().lower()


def fmt(v, numfmt="General"):
    """Wartosc komorki Excela -> tekst do pisma. Liczby wg formatu komorki:
    '0.00' -> 42,50; '#,##0.00' -> 1 234,50 (polski zapis)."""
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%d.%m.%Y") if v.time() == datetime.time() else v.strftime("%d.%m.%Y %H:%M")
    if isinstance(v, datetime.date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        m = re.search(r"\.(0+)", numfmt or "")
        if m or "#,##" in (numfmt or ""):
            s = "{:,.{}f}".format(v, len(m.group(1)) if m else 0)
            return s.replace(",", " ").replace(".", ",")
        if isinstance(v, float):
            return str(int(v)) if v.is_integer() else str(v).replace(".", ",")
    return str(v).strip()


def read_rows(xlsx):
    """[{pole: tekst}] z pierwszego arkusza; pierwszy wiersz = naglowki."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    rows = wb.worksheets[0].iter_rows()
    header = [key(c.value) if c.value is not None else "" for c in next(rows, [])]
    if not any(header):
        raise ValueError("Pierwszy wiersz arkusza musi zawierac nazwy pol")
    out = []
    for r in rows:
        if any(getattr(c, "value", None) not in (None, "") for c in r):
            row = dict.fromkeys((h for h in header if h), "")  # brakujace komorki = puste
            row.update({h: fmt(c.value, getattr(c, "number_format", "General"))
                        for h, c in zip(header, r) if h})
            out.append(row)
    wb.close()
    return header, out


# ------------------------------------------------------------------ DOCX ----


def _parts(doc):
    return [doc.part] + [r.target_part for r in doc.part.rels.values()
                         if r.reltype.endswith(("/header", "/footer", "/footnotes", "/endnotes"))]


def template_fields(doc):
    from docx.oxml.ns import qn
    found = []
    for part in _parts(doc):
        for p in part.element.iter(qn("w:p")):
            text = "".join(t.text or "" for t in p.iter(qn("w:t")))
            found += [key(m.group(1)) for m in FIELD.finditer(text)]
    return list(dict.fromkeys(found))


def _set_text(t, value):
    """Wpisuje tekst do w:t; znaki nowej linii z Excela -> w:br w tym samym runie."""
    from docx.oxml import OxmlElement
    lines = value.split("\n")
    t.text = lines[0]
    t.set(XML_SPACE, "preserve")
    for line in lines[1:]:
        br, nt = OxmlElement("w:br"), OxmlElement("w:t")
        nt.text = line
        nt.set(XML_SPACE, "preserve")
        t.addnext(br)
        br.addnext(nt)
        t = nt


def fill_paragraph(p, values):
    """Podmienia {pola} w akapicie. Word czesto tnie '{imie}' na kilka runow,
    wiec dziala na zlaczonym tekscie i odwzorowuje pozycje na runy.
    Zwraca nazwy pol nieznalezionych w danych."""
    from docx.oxml.ns import qn
    nodes = list(p.iter(qn("w:t")))
    texts = [t.text or "" for t in nodes]
    full = "".join(texts)
    missing = []
    starts, pos = [], 0
    for s in texts:
        starts.append(pos)
        pos += len(s)

    def node_at(i):
        n = 0
        while n + 1 < len(starts) and starts[n + 1] <= i:
            n += 1
        return n

    edits = {}  # nr wezla -> nowy tekst
    # od konca, zeby wczesniejsze pozycje sie nie przesuwaly
    for m in reversed(list(FIELD.finditer(full))):
        k = key(m.group(1))
        if k not in values:
            missing.append(k)
            continue
        a, b = m.start(), m.end()
        i, j = node_at(a), node_at(b - 1)
        ti = edits.get(i, texts[i])
        if i == j:
            edits[i] = ti[:a - starts[i]] + values[k] + ti[b - starts[i]:]
        else:
            tj = edits.get(j, texts[j])
            edits[j] = tj[b - starts[j]:]
            for n in range(i + 1, j):
                edits[n] = ""
            edits[i] = ti[:a - starts[i]] + values[k]
        texts = [edits.get(n, texts[n]) for n in range(len(texts))]
    for n, value in edits.items():
        _set_text(nodes[n], value)
    return missing


def fill_doc(template, values, dst):
    import docx
    from docx.oxml.ns import qn
    doc = docx.Document(template)
    missing = []
    for part in _parts(doc):
        for p in list(part.element.iter(qn("w:p"))):
            missing += fill_paragraph(p, values)
    doc.save(dst)
    return missing


# ------------------------------------------------------------------- PDF ----


class Word:
    """Konwersja DOCX -> PDF przez zainstalowanego Worda (jedno uruchomienie na cala serie)."""

    def __enter__(self):
        try:
            import win32com.client
            self.app = win32com.client.DispatchEx("Word.Application")
        except Exception:
            raise ValueError("Nie udalo sie uruchomic Microsoft Word - odznacz opcje PDF "
                             "albo zainstaluj Worda (i biblioteke pywin32: install.bat)")
        self.app.Visible = False
        self.app.DisplayAlerts = 0
        return self

    def to_pdf(self, docx_path, pdf_path):
        d = self.app.Documents.Open(os.path.abspath(docx_path), ReadOnly=True)
        try:
            d.SaveAs2(os.path.abspath(pdf_path), FileFormat=17)  # 17 = PDF
        finally:
            d.Close(False)

    def __exit__(self, *exc):
        self.app.Quit()


# ------------------------------------------------------------------ wsad ----

BAD_CHARS = re.compile(r'[\\/:*?"<>|\r\n\t]+')


def file_name(pattern, values, nr):
    vals = dict(values, nr=str(nr))
    name = FIELD.sub(lambda m: vals.get(key(m.group(1)), ""), pattern)
    name = BAD_CHARS.sub("_", name).strip(" ._")
    return name[:120] or "pismo_%d" % nr


def run(template, xlsx, out_dir, pattern="pismo_{nr}", pdf=False, log=print):
    import docx
    header, rows = read_rows(xlsx)
    fields = template_fields(docx.Document(template))
    log("Pola w szablonie: %s" % (", ".join(fields) or "brak"))
    log("Kolumny w arkuszu: %s" % ", ".join(h for h in header if h))
    unknown = [f for f in fields if f not in header and f != "nr"]
    if unknown:
        log("UWAGA: brak kolumn dla pol: %s - zostana w pismach bez zmian" % ", ".join(unknown))
    if not rows:
        log("Arkusz nie zawiera danych.")
        return []
    os.makedirs(out_dir, exist_ok=True)
    used, res = set(), []
    word = Word().__enter__() if pdf else None
    try:
        for nr, values in enumerate(rows, 1):
            values = dict(values, nr=values.get("nr") or str(nr))
            name = file_name(pattern, values, nr)
            base, k = name, 2
            while name.lower() in used:  # te same nazwiska -> _2, _3...
                name, k = "%s_%d" % (base, k), k + 1
            used.add(name.lower())
            dst = os.path.join(out_dir, name + ".docx")
            try:
                fill_doc(template, values, dst)
                if word:
                    word.to_pdf(dst, os.path.join(out_dir, name + ".pdf"))
                res.append(dst)
                log("  %d/%d OK -> %s%s" % (nr, len(rows), name, " (.docx + .pdf)" if word else ".docx"))
            except Exception:
                log("  %d/%d BLAD: %s\n%s" % (nr, len(rows), name, traceback.format_exc()))
    finally:
        if word:
            word.__exit__(None, None, None)
    log("Zakonczono: %d z %d pism." % (len(res), len(rows)))
    return res


# -------------------------------------------------------------------- GUI ----


def gui():
    import tkinter as tk
    from tkinter import filedialog, ttk, scrolledtext

    root = tk.Tk()
    root.title("Korespondencja seryjna (DOCX + XLSX)")
    root.geometry("800x540")
    pad = dict(padx=6, pady=3)
    ex = os.path.join(APP_DIR, "przyklad")

    v_tpl = tk.StringVar(value=os.path.join(ex, "szablon.docx"))
    v_xls = tk.StringVar(value=os.path.join(ex, "dane.xlsx"))
    v_out = tk.StringVar(value=os.path.join(APP_DIR, "OUTPUT"))
    v_pat = tk.StringVar(value="{nr}_{nazwisko}_{imie}")
    v_pdf = tk.BooleanVar(value=False)

    f = ttk.Frame(root)
    f.pack(fill="x", **pad)
    rows = [("Szablon DOCX:", v_tpl, [("Word", "*.docx")]),
            ("Dane XLSX:", v_xls, [("Excel", "*.xlsx")]),
            ("Folder wyjsciowy:", v_out, None)]
    for i, (lab, var, types) in enumerate(rows):
        ttk.Label(f, text=lab).grid(row=i, column=0, sticky="w", **pad)
        ttk.Entry(f, textvariable=var, width=64).grid(row=i, column=1, **pad)
        cmd = (lambda v=var, t=types: v.set(filedialog.askopenfilename(filetypes=t) or v.get())) if types \
            else (lambda v=var: v.set(filedialog.askdirectory() or v.get()))
        ttk.Button(f, text="Wybierz...", command=cmd).grid(row=i, column=2, **pad)
    ttk.Label(f, text="Nazwa pliku:").grid(row=3, column=0, sticky="w", **pad)
    ttk.Entry(f, textvariable=v_pat, width=64).grid(row=3, column=1, **pad)
    ttk.Label(f, text="pola z arkusza w {klamrach}; {nr} = numer wiersza",
              foreground="#666").grid(row=4, column=1, sticky="w", padx=6)
    ttk.Checkbutton(f, text="Utworz tez PDF (wymaga zainstalowanego Microsoft Word)",
                    variable=v_pdf).grid(row=5, column=1, sticky="w", **pad)

    log_box = scrolledtext.ScrolledText(root, height=16)
    log_box.pack(fill="both", expand=True, **pad)

    def log(msg):
        def put():
            log_box.insert("end", str(msg) + "\n")
            log_box.see("end")
        root.after(0, put)

    btn = ttk.Button(root, text="Generuj pisma")
    btn.pack(pady=6)

    def start():
        tpl, xls, out = (v.get().strip('" ') for v in (v_tpl, v_xls, v_out))
        if not os.path.isfile(tpl) or not os.path.isfile(xls):
            return log("Wskaz istniejacy szablon DOCX i plik XLSX.")
        if not out:
            return log("Wskaz folder wyjsciowy.")
        btn.config(state="disabled")
        log_box.delete("1.0", "end")

        def work():
            try:
                if v_pdf.get():
                    import pythoncom  # COM w watku innym niz glowny
                    pythoncom.CoInitialize()
                run(tpl, xls, out, v_pat.get().strip() or "pismo_{nr}", v_pdf.get(), log)
            except ValueError as e:
                log("BLAD: %s" % e)
            except Exception:
                log("BLAD:\n" + traceback.format_exc())
            finally:
                root.after(0, lambda: btn.config(state="normal"))

        threading.Thread(target=work, daemon=True).start()

    btn.config(command=start)
    if "--selftest" in sys.argv:
        root.after(200, root.destroy)
    import aktualizacja
    aktualizacja.start(root, "DawidBochno/Korespondencja-seryjna", "main", "korespondencja.py")
    root.mainloop()


# --------------------------------------------------------------- selftest ----


def selftest(pdf=False):
    import tempfile
    import docx
    from openpyxl import Workbook

    assert fmt(1234.0) == "1234" and fmt(12.5) == "12,5"
    assert fmt(42.5, "0.00") == "42,50" and fmt(1234.5, "#,##0.00") == "1 234,50"
    assert fmt(7, "0.00") == "7,00" and fmt(1500, "#,##0") == "1 500"
    assert fmt(datetime.datetime(2026, 10, 1)) == "01.10.2026"
    assert file_name("{nr}_{Nazwisko}/{x}", {"nazwisko": "Nowak: A"}, 3) == "3_Nowak_ A"
    assert file_name("{x}", {}, 7) == "pismo_7"

    tmp = tempfile.mkdtemp()
    tpl = os.path.join(tmp, "s.docx")
    d = docx.Document()
    p = d.add_paragraph("Szanowny Panie ")
    p.add_run("{Imi")              # pole pociete na 3 runy, jak robi Word
    p.add_run("e").bold = True
    p.add_run("} {nazwisko}, kwota {kwota} zl, {nieznane}.")
    d.add_paragraph("Adres: {adres}")
    d.sections[0].header.paragraphs[0].text = "Znak: {nr}"
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Termin: {termin}"
    d.save(tpl)

    xls = os.path.join(tmp, "d.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(["Imie", "Nazwisko", "Kwota", "Adres", "Termin"])
    ws.append(["Jan", "Kowalski", 1250.5, "ul. Polna 1\n00-001 Warszawa", datetime.datetime(2026, 11, 15)])
    ws.append([None, None, None, None, None])  # pusty wiersz - pomijany
    ws.append(["Anna", "Kowalski", 300, "ul. Lesna 2", None])
    wb.save(xls)

    assert template_fields(docx.Document(tpl)) == ["imie", "nazwisko", "kwota", "nieznane",
                                                   "adres", "termin", "nr"]
    msgs = []
    res = run(tpl, xls, os.path.join(tmp, "out"), "{nazwisko}", pdf, msgs.append)
    assert [os.path.basename(r) for r in res] == ["Kowalski.docx", "Kowalski_2.docx"], res
    assert any("brak kolumn dla pol: nieznane" in m for m in msgs), msgs
    o = docx.Document(res[0])
    assert o.paragraphs[0].text == "Szanowny Panie Jan Kowalski, kwota 1250,5 zl, {nieznane}.", o.paragraphs[0].text
    assert o.paragraphs[0].runs[1].text == "Jan"  # wartosc w runie, w ktorym zaczyna sie pole
    assert o.paragraphs[1].text == "Adres: ul. Polna 1\n00-001 Warszawa", repr(o.paragraphs[1].text)
    assert o.sections[0].header.paragraphs[0].text == "Znak: 1"
    assert o.tables[0].cell(0, 0).text == "Termin: 15.11.2026"
    o2 = docx.Document(res[1])
    assert o2.sections[0].header.paragraphs[0].text == "Znak: 2"
    assert o2.tables[0].cell(0, 0).text == "Termin: "
    if pdf:
        assert os.path.isfile(res[0][:-5] + ".pdf")
    import aktualizacja
    aktualizacja.selftest()
    print("selftest OK" + (" (z PDF)" if pdf else ""))


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest(pdf="--pdf" in sys.argv)
        if "--gui" in sys.argv:
            gui()
    else:
        gui()
